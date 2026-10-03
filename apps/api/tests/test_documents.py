from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import func, select

from app.api.dependencies import get_database_session
from app.db.models import AuditEvent, Document, Project, ProjectMember
from app.domain.document_parse import MAX_BYTES, parse_document
from app.domain.manual_state import StateError
from app.main import app
from app.repositories.audit import AuditRepository
from app.services.demo_seed import DEMO_PROJECT_ID, seed_demo
from app.services.project_access import LOCAL_DEMO_ACTOR_ID

ROOT = f"/api/v1/projects/{DEMO_PROJECT_ID}/documents"


def pdf_bytes():
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 30 220 Td (Launch approval is pending.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    result = BytesIO()
    writer.write(result)
    return result.getvalue()


def docx_bytes():
    doc = DocxDocument()
    doc.add_heading("Scope", 1)
    doc.add_paragraph("SSO launches in phase 2.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Approval pending"
    result = BytesIO()
    doc.save(result)
    return result.getvalue()


@pytest.mark.parametrize(
    "filename,raw,page,section",
    [
        ("scope.txt", b"SSO launches in phase 2.", None, "Document"),
        ("scope.md", b"# Scope\nSSO launches in phase 2.", None, "Scope"),
        ("scope.docx", docx_bytes(), None, "Scope"),
        ("scope.pdf", pdf_bytes(), 1, None),
    ],
)
def test_parsers_preserve_text_and_location(filename, raw, page, section):
    name, _, segments = parse_document(filename, raw)
    assert name == filename and segments[0]["page"] == page and segments[0]["section"] == section
    assert segments[0]["text"].strip()


@pytest.mark.parametrize(
    "filename,raw",
    [
        ("x.exe", b"hello"),
        ("x.txt", b""),
        ("x.txt", b"  \n"),
        ("x.txt", b"\x00"),
        ("x.txt", b"\xff"),
        ("x.docx", b"PKbroken"),
        ("x.pdf", b"%PDF-broken"),
        ("x.pdf", b"plain text"),
        ("x.txt", b"x" * (MAX_BYTES + 1)),
        ("x.txt", b"x" * 200001),
    ],
)
def test_invalid_document_rejected(filename, raw):
    with pytest.raises(StateError):
        parse_document(filename, raw)


def test_zip_bomb_rejected():
    raw = BytesIO()
    with ZipFile(raw, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 100000)
    with pytest.raises(StateError):
        parse_document("x.docx", raw.getvalue())


def test_image_only_and_encrypted_pdf_rejected():
    for encrypted in [False, True]:
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        if encrypted:
            writer.encrypt("password")
        raw = BytesIO()
        writer.write(raw)
        with pytest.raises(StateError):
            parse_document("x.pdf", raw.getvalue())


@pytest.fixture
def client(db_session):
    seed_demo(db_session)
    app.dependency_overrides[get_database_session] = lambda: db_session
    try:
        with TestClient(app) as client:
            yield client
    finally:
        app.dependency_overrides.pop(get_database_session, None)


@pytest.mark.integration
@pytest.mark.parametrize(
    "filename,raw",
    [
        ("x.txt", b"API scope"),
        ("x.md", b"# Scope\nAPI scope"),
        ("x.pdf", pdf_bytes()),
        ("x.docx", docx_bytes()),
    ],
)
def test_upload_duplicate_read_delete_without_changing_truth(db_session, client, filename, raw):
    before = client.get(f"/api/v1/projects/{DEMO_PROJECT_ID}/workspace").json()
    result = client.post(ROOT, params={"filename": filename}, content=raw)
    assert result.status_code == 201, result.text
    doc = result.json()
    assert not doc["duplicate"] and doc["projectId"] == str(DEMO_PROJECT_ID)
    duplicate = client.post(ROOT, params={"filename": "renamed.txt"}, content=raw).json()
    assert duplicate["id"] == doc["id"] and duplicate["duplicate"]
    assert len(client.get(ROOT).json()) == 1
    detail = client.get(f"{ROOT}/{doc['id']}").json()
    assert detail["segments"][0]["text"].strip() and "raw_bytes" not in detail
    after = client.get(f"/api/v1/projects/{DEMO_PROJECT_ID}/workspace").json()
    for category in [
        "requirements",
        "decisions",
        "milestones",
        "risks",
        "commitments",
        "dependencies",
    ]:
        assert before[category] == after[category]
    assert client.delete(f"{ROOT}/{doc['id']}").status_code == 204
    assert client.get(f"{ROOT}/{doc['id']}").status_code == 404
    assert client.get(ROOT).json() == []
    audits = db_session.scalars(select(AuditEvent).where(AuditEvent.entity_id == doc["id"])).all()
    assert len(audits) == 2 and all(x.actor_id == LOCAL_DEMO_ACTOR_ID for x in audits)


@pytest.mark.integration
@pytest.mark.parametrize("method", ["get", "post", "delete"])
def test_document_access_does_not_leak_other_projects(db_session, client, method):
    doc = client.post(ROOT, params={"filename": "x.txt"}, content=b"scoped evidence").json()
    hidden = Project(name="Hidden")
    db_session.add(hidden)
    db_session.flush()
    path = f"/api/v1/projects/{hidden.id}/documents" + ("" if method == "post" else f"/{doc['id']}")
    response = getattr(client, method)(
        path,
        **({"params": {"filename": "x.txt"}, "content": b"outside"} if method == "post" else {}),
    )
    assert response.status_code == 404
    db_session.add(
        ProjectMember(project_id=hidden.id, user_id=LOCAL_DEMO_ACTOR_ID, role="PRODUCT_OWNER")
    )
    db_session.flush()
    if method != "post":
        assert getattr(client, method)(path).status_code == 404


@pytest.mark.integration
def test_oversized_stream_and_parse_errors_write_nothing(db_session, client):
    for name, raw in [("x.txt", b"x" * (MAX_BYTES + 1)), ("x.pdf", b"bad")]:
        assert client.post(ROOT, params={"filename": name}, content=raw).status_code in [413, 422]
    assert db_session.scalar(select(func.count()).select_from(Document)) == 0


@pytest.mark.integration
@pytest.mark.parametrize("remove", [False, True])
def test_document_mutation_rolls_back_if_audit_fails(db_session, client, monkeypatch, remove):
    doc = (
        client.post(ROOT, params={"filename": "x.txt"}, content=b"existing").json()
        if remove
        else None
    )

    def fail(*args, **kwargs):
        raise RuntimeError("audit failure")

    monkeypatch.setattr(AuditRepository, "append", fail)
    with pytest.raises(RuntimeError):
        if remove:
            client.delete(f"{ROOT}/{doc['id']}")
        else:
            client.post(ROOT, params={"filename": "x.txt"}, content=b"new")
    assert db_session.scalar(select(func.count()).select_from(Document)) == (1 if remove else 0)
