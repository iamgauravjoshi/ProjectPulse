"""Bounded, local open-source parsing. No AI or canonical-state writes."""

from io import BytesIO
from pathlib import PurePosixPath
from typing import Any
from zipfile import ZipFile

from docx import Document as DocxDocument
from pypdf import PdfReader, apply_configuration

from app.domain.manual_state import StateError

MAX_BYTES = 5 * 1024 * 1024
MAX_TEXT = 200_000
# Bound each compressed PDF stream as well as the uploaded file.


def parse_document(filename: str, raw: bytes) -> tuple[str, str, list[dict[str, Any]]]:
    name = filename.replace("\\", "/").rsplit("/", 1)[-1].strip()
    if not name or len(name) > 240 or any(ord(c) < 32 for c in name):
        raise StateError("INVALID_FILENAME", "Choose a valid filename.")
    extension = PurePosixPath(name).suffix.lower().lstrip(".")
    if extension not in {"pdf", "docx", "txt", "md", "markdown"}:
        raise StateError("UNSUPPORTED_DOCUMENT", "Choose a PDF, DOCX, TXT or Markdown file.")
    if not raw or len(raw) > MAX_BYTES:
        raise StateError(
            "DOCUMENT_SIZE", "Files must contain data and be no larger than 5 MiB.", 413
        )
    segments: list[dict[str, Any]] = []
    try:
        if extension == "pdf":
            if not raw.startswith(b"%PDF-"):
                raise ValueError("PDF signature")
            with apply_configuration(
                zlib_maximum_output_length=4 * 1024 * 1024,
                maximum_declared_stream_length=4 * 1024 * 1024,
                page_tree_maximum_entries=1000,
                xform_maximum_invocations_per_extraction=100,
                jbig2dec_binary=None,
            ):
                pdf = PdfReader(BytesIO(raw), strict=True)
                if pdf.is_encrypted or len(pdf.pages) > 200:
                    raise ValueError("Encrypted or too many pages")
                for index, page in enumerate(pdf.pages):
                    value = page.extract_text() or ""
                    segments.append({"page": index + 1, "section": None, "text": value})
                    if sum(len(x["text"]) for x in segments) > MAX_TEXT:
                        raise ValueError("Text limit")
        elif extension == "docx":
            with ZipFile(BytesIO(raw)) as archive:
                entries = archive.infolist()
                if (
                    len(entries) > 1000
                    or sum(x.file_size for x in entries) > 20 * 1024 * 1024
                    or any(
                        x.flag_bits & 1
                        or x.file_size > 5 * 1024 * 1024
                        or x.file_size > max(1, x.compress_size) * 100
                        for x in entries
                    )
                ):
                    raise ValueError("Archive bounds")
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("DOCX structure")
            document = DocxDocument(BytesIO(raw))
            section = "Document"
            for paragraph in document.paragraphs:
                if (
                    paragraph.style is not None
                    and paragraph.style.name.startswith("Heading")
                    and paragraph.text.strip()
                ):
                    section = paragraph.text[:240]
                if paragraph.text.strip():
                    segments.append({"page": None, "section": section, "text": paragraph.text})
            for table in document.tables:
                for row in table.rows:
                    segments.append(
                        {
                            "page": None,
                            "section": "Table",
                            "text": " | ".join(cell.text for cell in row.cells),
                        }
                    )
        else:
            value = raw.decode("utf-8-sig")
            if "\x00" in value:
                raise ValueError("Binary text")
            if extension in {"md", "markdown"}:
                section = "Document"
                lines: list[str] = []
                for line in value.splitlines():
                    if line.startswith("#"):
                        if lines:
                            segments.append(
                                {"page": None, "section": section, "text": "\n".join(lines)}
                            )
                        section = line.lstrip("# ")[:240] or "Document"
                        lines = []
                    lines.append(line)
                if lines:
                    segments.append({"page": None, "section": section, "text": "\n".join(lines)})
            else:
                segments = [{"page": None, "section": "Document", "text": value}]
        segments = [x for x in segments if x["text"].strip()]
        if not segments or sum(len(x["text"]) for x in segments) > MAX_TEXT:
            raise ValueError("Empty or too much text")
    except Exception:
        raise StateError(
            "DOCUMENT_PARSE_FAILED",
            "Unreadable, malformed, encrypted or oversized document. "
            "Scanned PDFs need OCR before upload.",
        ) from None
    return name, extension, segments
